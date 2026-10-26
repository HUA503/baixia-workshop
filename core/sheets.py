from __future__ import annotations

import csv
import io
import json
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .utils import unique_path


def _norm(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if hasattr(value, "isoformat"):
        try:
            return value.isoformat(sep=" ")
        except TypeError:
            return value.isoformat()
    return value


NUMBER_RE = re.compile(r"^-?(?:0|[1-9]\d{0,9})(?:\.\d{1,6})?$")


def _excel_value(value):
    """CSV cells are strings; turn plain numeric cells into numbers for Excel."""
    if not isinstance(value, str):
        return value
    text = value.strip()
    if not text or not NUMBER_RE.match(text):
        return value
    try:
        return float(text) if "." in text else int(text)
    except ValueError:
        return value


def _read_csv(path: Path) -> list[list]:
    raw = path.read_bytes()
    text = None
    for enc in ("utf-8-sig", "utf-8", "gb18030", "gbk", "big5", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        text = raw.decode("utf-8", errors="replace")
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except Exception:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    return [[_norm(c) for c in row] for row in reader]


def _read_workbook(path: Path) -> list[tuple[str, list[list]]]:
    if path.suffix.lower() == ".csv":
        return [(path.stem, _read_csv(path))]
    wb = load_workbook(path, data_only=True, read_only=True)
    sheets = []
    try:
        for ws in wb.worksheets:
            rows = []
            for row in ws.iter_rows(values_only=True):
                rows.append([_norm(v) for v in row])
            while rows and all(c == "" for c in rows[-1]):
                rows.pop()
            sheets.append((ws.title or "Sheet", rows))
    finally:
        wb.close()
    return sheets


def _as_text(value) -> str:
    if isinstance(value, str):
        return value.replace("|", "\\|").replace("\n", "<br>")
    return str(value).replace("|", "\\|").replace("\n", "<br>")


def _md_table(rows: list[list], header: bool) -> str:
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [(r + [""] * width)[:width] for r in rows]
    if header:
        head, body = rows[0], rows[1:]
    else:
        head = [f"列{i + 1}" for i in range(width)]
        body = rows
    lines = ["| " + " | ".join(_as_text(c) for c in head) + " |",
             "| " + " | ".join(["---"] * width) + " |"]
    for row in body:
        lines.append("| " + " | ".join(_as_text(c) for c in row) + " |")
    return "\n".join(lines)


def _json_rows(rows: list[list], header: bool, sheet_name: str):
    if not rows:
        return []
    width = max(len(r) for r in rows)
    rows = [(r + [""] * width)[:width] for r in rows]
    if not header:
        return rows
    keys = [str(c) if str(c) else f"列{i + 1}" for i, c in enumerate(rows[0])]
    seen: dict[str, int] = {}
    for i, key in enumerate(keys):
        if key in seen:
            seen[key] += 1
            keys[i] = f"{key}_{seen[key]}"
        else:
            seen[key] = 0
    result = []
    for row in rows[1:]:
        result.append({keys[i]: row[i] for i in range(width)})
    return result


def _write_csv(rows: list[list], dest: Path) -> None:
    with dest.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh)
        for row in rows:
            writer.writerow(row)


def _sheet_filename(stem: str, sheet: str, ext: str) -> str:
    safe = "".join(ch if ch not in '<>:"/\\|?*' else "_" for ch in sheet).strip() or "sheet"
    return f"{stem}-{safe}{ext}"


def xlsx_convert(job, inputs: list[Path], opts: dict):
    src = inputs[0]
    fmt = opts.get("format", "md")
    all_sheets = opts.get("sheet", "first") == "all"
    header = bool(opts.get("header", True))
    job.report(10, "读取表格")
    sheets = _read_workbook(src)
    if not all_sheets:
        sheets = sheets[:1]
    if not sheets:
        raise RuntimeError("表格里没有内容")
    stem = src.stem
    if fmt == "xlsx":
        if src.suffix.lower() in {".xlsx", ".xlsm"}:
            import shutil

            dest = unique_path(job.outdir, f"{stem}.xlsx")
            shutil.copy2(src, dest)
            job.add_output(dest)
            return [dest]
        wb = Workbook()
        wb.remove(wb.active)
        for name, rows in sheets:
            ws = wb.create_sheet(title=(name or "Sheet")[:31])
            for row in rows:
                ws.append([_excel_value(value) for value in row])
        dest = unique_path(job.outdir, f"{stem}.xlsx")
        wb.save(dest)
        wb.close()
        job.add_output(dest)
        return [dest]
    if fmt == "csv":
        outputs = []
        for name, rows in sheets:
            target = unique_path(job.outdir, _sheet_filename(stem, name, ".csv"))
            _write_csv(rows, target)
            outputs.append(target)
            job.add_output(target)
        return outputs
    if fmt == "json":
        if all_sheets:
            data = {}
            for name, rows in sheets:
                data[name] = _json_rows(rows, header, name)
        else:
            data = _json_rows(sheets[0][1], header, sheets[0][0])
        dest = unique_path(job.outdir, f"{stem}.json")
        dest.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        job.add_output(dest)
        return [dest]
    parts = []
    for name, rows in sheets:
        if all_sheets:
            parts.append(f"## {name}\n")
        parts.append(_md_table(rows, header))
        parts.append("")
    dest = unique_path(job.outdir, f"{stem}.md")
    dest.write_text("\n".join(parts), encoding="utf-8")
    job.add_output(dest)
    return [dest]
