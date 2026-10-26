from __future__ import annotations

import base64
import gc
import html as html_lib
import mimetypes
import os
import re
import shutil
import subprocess
import threading
import time
from pathlib import Path

import pymupdf
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

from .utils import unique_path

DOC_CSS = """
body { font-family: sans-serif; font-size: 10.5pt; line-height: 1.6; color: #141414; }
h1 { font-size: 22pt; margin: 18pt 0 8pt; }
h2 { font-size: 16pt; margin: 14pt 0 6pt; }
h3 { font-size: 13pt; margin: 12pt 0 5pt; }
h4, h5, h6 { font-size: 11.5pt; margin: 10pt 0 4pt; }
p { margin: 0 0 7pt; }
ul, ol { margin: 0 0 8pt 18pt; padding: 0; }
li { margin: 0 0 3pt; }
blockquote { margin: 6pt 0 8pt 14pt; padding-left: 8pt; border-left: 2pt solid #cccccc; color: #444444; }
pre { font-family: monospace; font-size: 9pt; background: #f5f5f5; padding: 6pt; margin: 6pt 0 8pt; white-space: pre-wrap; }
code { font-family: monospace; font-size: 9pt; }
table { border-collapse: collapse; width: 100%; margin: 6pt 0 10pt; font-size: 9pt; }
th, td { border: 0.5pt solid #999999; padding: 3pt 5pt; text-align: left; vertical-align: top; }
th { background: #f0f0f0; font-weight: bold; }
img { max-width: 100%; }
hr { border: none; border-top: 0.5pt solid #cccccc; margin: 10pt 0; }
a { color: #1a5fb4; }
.small { font-size: 9pt; color: #666666; }
"""


# ───────────────────────── markdown → html ─────────────────────────

def _md_inline(text: str) -> str:
    parts = re.split(r"(`[^`]+`|!\[[^\]]*\]\([^)]*\)|\[[^\]]+\]\([^)]*\))", text)
    out: list[str] = []
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`") and len(part) >= 2:
            out.append(f"<code>{html_lib.escape(part[1:-1])}</code>")
        elif part.startswith("!["):
            m = re.match(r"!\[([^\]]*)\]\(([^)]*)\)", part)
            if m:
                alt = html_lib.escape(m.group(1))
                src = _safe_src(m.group(2))
                if src:
                    out.append(f'<img src="{html_lib.escape(src, quote=True)}" alt="{alt}">')
        elif part.startswith("["):
            m = re.match(r"\[([^\]]+)\]\(([^)]*)\)", part)
            if m:
                label = html_lib.escape(m.group(1))
                href = _safe_src(m.group(2))
                if href:
                    out.append(f'<a href="{html_lib.escape(href, quote=True)}">{label}</a>')
                else:
                    out.append(label)
        else:
            s = html_lib.escape(part)
            s = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", s, flags=re.S)
            s = re.sub(r"__(.+?)__", r"<strong>\1</strong>", s, flags=re.S)
            s = re.sub(r"~~(.+?)~~", r"<del>\1</del>", s, flags=re.S)
            s = re.sub(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)", r"<em>\1</em>", s, flags=re.S)
            s = re.sub(r"(?<!_)_(?!_)(.+?)(?<!_)_(?!_)", r"<em>\1</em>", s, flags=re.S)
            out.append(s)
    return "".join(out)


def _safe_src(src: str) -> str:
    src = (src or "").strip().strip("<>")
    if re.match(r"^(https?://|data:image/|assets/|\.?/|[A-Za-z0-9_\-]+\.(png|jpe?g|gif|webp|svg|avif))", src, re.I):
        return src
    return ""


def _md_table(lines: list[str], i: int) -> tuple[str, int]:
    rows = []
    while i < len(lines) and "|" in lines[i] and lines[i].strip():
        row = lines[i].strip()
        if row.startswith("|"):
            row = row[1:]
        if row.endswith("|"):
            row = row[:-1]
        rows.append([c.strip() for c in row.split("|")])
        i += 1
    if len(rows) >= 2 and re.match(r"^:?-{2,}:?$", rows[1][0].replace(" ", "")):
        header, body = rows[0], rows[2:]
    else:
        header, body = rows[0], rows[1:]
    width = len(header)
    frag = ["<table>", "<thead><tr>"]
    for cell in header:
        frag.append(f"<th>{_md_inline(cell)}</th>")
    frag.append("</tr></thead><tbody>")
    for row in body:
        frag.append("<tr>")
        for cell in (row + [""] * width)[:width]:
            frag.append(f"<td>{_md_inline(cell)}</td>")
        frag.append("</tr>")
    frag.append("</tbody></table>")
    return "".join(frag), i


def md_to_html(md: str, title: str = "", full: bool = True) -> str:
    lines = (md or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")
    out: list[str] = []
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith("```") or stripped.startswith("~~~"):
            fence = stripped[:3]
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith(fence):
                code_lines.append(lines[i])
                i += 1
            i += 1
            out.append(f"<pre><code>{html_lib.escape(chr(10).join(code_lines))}</code></pre>")
            continue
        if not stripped:
            i += 1
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            out.append(f"<h{level}>{_md_inline(m.group(2))}</h{level}>")
            i += 1
            continue
        if re.match(r"^(\*\s*){3,}$|^(-\s*){3,}$|^(_\s*){3,}$", stripped):
            out.append("<hr>")
            i += 1
            continue
        if stripped.startswith(">"):
            quote = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip()[1:].lstrip())
                i += 1
            out.append(f"<blockquote><p>{_md_inline(' '.join(quote))}</p></blockquote>")
            continue
        if "|" in stripped and i + 1 < len(lines) and re.search(r"\|?\s*:?-{2,}", lines[i + 1]):
            table, i = _md_table(lines, i)
            out.append(table)
            continue
        lm = re.match(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$", line)
        if lm:
            ordered = lm.group(2)[0].isdigit()
            tag = "ol" if ordered else "ul"
            items = []
            while i < len(lines):
                im = re.match(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$", lines[i])
                if not im:
                    break
                items.append(_md_inline(im.group(3)))
                i += 1
            out.append(f"<{tag}>" + "".join(f"<li>{x}</li>" for x in items) + f"</{tag}>")
            continue
        para = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i]
            ns = nxt.strip()
            if (not ns or ns.startswith(("#", "```", "~~~", ">")) or re.match(r"^(\s*)([-*+]|\d+[.)])\s+", nxt)
                    or re.match(r"^(\*\s*){3,}$|^(-\s*){3,}$", ns)):
                break
            para.append(ns)
            i += 1
        out.append(f"<p>{_md_inline(' '.join(para))}</p>")
    body = "\n".join(out)
    if not full:
        return body
    return (
        "<!DOCTYPE html><html lang=\"zh-CN\"><head><meta charset=\"utf-8\">"
        f"<meta name=\"viewport\" content=\"width=device-width, initial-scale=1\"><title>{html_lib.escape(title or '文档')}</title>"
        "<style>"
        "body{margin:0;background:#f6f6f4;color:#141414;font-family:-apple-system,'Segoe UI','Microsoft YaHei',sans-serif;}"
        ".wrap{max-width:820px;margin:0 auto;padding:56px 28px;}"
        ".doc{background:#fff;border:1px solid #e5e5e0;border-radius:4px;padding:56px 64px;box-shadow:0 1px 2px rgba(0,0,0,.04);}"
        + DOC_CSS +
        "pre{background:#f7f7f5;border:1px solid #ecece8;border-radius:2px;padding:12px;overflow:auto;}"
        "th,td{border-color:#deded8;}th{background:#fafaf7;}"
        "@media(max-width:640px){.wrap{padding:20px 12px}.doc{padding:28px 20px}}"
        "</style></head><body><div class=\"wrap\"><article class=\"doc\">"
        f"{body}</article></div></body></html>"
    )


def html_to_pdf(html: str, dest: Path, css: str = DOC_CSS, archive_dir: Path | None = None) -> Path:
    archive = pymupdf.Archive(str(archive_dir)) if archive_dir and archive_dir.exists() else None
    story = pymupdf.Story(html=html, user_css=css, archive=archive)
    writer = pymupdf.DocumentWriter(str(dest))
    page_rect = pymupdf.paper_rect("a4")
    margin = 52
    area = page_rect + (margin, margin, -margin, -margin)
    more = 1
    pages = 0
    while more and pages < 800:
        device = writer.begin_page(page_rect)
        more, _ = story.place(area)
        story.draw(device)
        writer.end_page()
        pages += 1
    writer.close()
    del writer
    gc.collect()
    doc = None
    try:
        doc = pymupdf.open(str(dest))
        doc.subset_fonts()
        tmp = dest.with_name(dest.stem + ".subset.pdf")
        doc.save(str(tmp), garbage=4, deflate=True, deflate_images=True, deflate_fonts=True)
        doc.close()
        gc.collect()
        dest.unlink(missing_ok=True)
        tmp.replace(dest)
    except Exception:
        if doc is not None:
            try:
                doc.close()
            except Exception:
                pass
    return dest


# ───────────────────────── docx → markdown ─────────────────────────

def _iter_blocks(doc):
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def _local_name(style_name: str) -> str:
    return (style_name or "").strip()


def _heading_level(paragraph: Paragraph) -> int | None:
    name = _local_name(paragraph.style.name if paragraph.style is not None else "")
    m = re.match(r"^(?:Heading|heading)\s*(\d)$", name)
    if m:
        return int(m.group(1))
    m = re.match(r"^标题\s*(\d)$", name)
    if m:
        return int(m.group(1))
    if name in {"Title", "标题"}:
        return 1
    if name in {"Subtitle", "副标题"}:
        return 2
    return None


def _list_kind(paragraph: Paragraph) -> tuple[str, int]:
    name = (_local_name(paragraph.style.name if paragraph.style is not None else "")).lower()
    level = 0
    ppr = paragraph._p.pPr
    if ppr is not None and ppr.numPr is not None:
        numpr = ppr.numPr
        ilvl = numpr.ilvl.val if numpr.ilvl is not None else 0
        level = min(3, int(ilvl or 0))
    if "list" in name or "列表" in name:
        ordered = any(k in name for k in ("number", "numbered", "编号", "数字"))
        return ("ol" if ordered else "ul"), level
    if ppr is not None and ppr.numPr is not None:
        return "ul", level
    return "", level


def _wrap_marker(text: str, marker: str) -> str:
    m = re.match(r"^(\s*)(.*?)(\s*)$", text, flags=re.S)
    if not m or not m.group(2):
        return text
    return f"{m.group(1)}{marker}{m.group(2)}{marker}{m.group(3)}"


def _run_md(run) -> str:
    text = run.text or ""
    if not text:
        return ""
    if not text.strip():
        return text
    name = (run.font.name or "").lower()
    if any(k in name for k in ("consolas", "courier", "mono")):
        text = _wrap_marker(text, "`")
    bold = bool(run.bold)
    italic = bool(run.italic)
    if bold and italic:
        text = _wrap_marker(text, "***")
    elif bold:
        text = _wrap_marker(text, "**")
    elif italic:
        text = _wrap_marker(text, "*")
    if run.font.strike:
        text = _wrap_marker(text, "~~")
    return text


def _para_md(paragraph: Paragraph, image_map: dict, images_mode: str) -> str:
    pieces: list[str] = []
    try:
        inner = list(paragraph.iter_inner_content())
    except Exception:
        inner = list(paragraph.runs)
    for item in inner:
        cls = item.__class__.__name__
        if cls == "Hyperlink":
            label = "".join(_run_md(r) for r in getattr(item, "runs", [])) or (item.text or "")
            url = getattr(item, "url", "") or ""
            if url:
                pieces.append(f"[{label or url}]({url})")
            else:
                pieces.append(label)
        else:
            pieces.append(_run_md(item))
    text = "".join(pieces).replace("\t", "    ").strip()
    name = (_local_name(paragraph.style.name if paragraph.style is not None else "")).lower()
    if text and ("quote" in name or "引用" in name):
        text = "> " + text
    images = []
    if images_mode != "none":
        for blip in paragraph._p.findall(".//" + qn("a:blip")):
            rid = blip.get(qn("r:embed"))
            if rid and rid in image_map:
                images.append(f"![]({image_map[rid]})")
    if images:
        text = (text + "\n\n" if text else "") + "\n\n".join(images)
    return text


def _table_md(table: Table) -> str:
    rows: list[list[str]] = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            value = " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip())
            value = value.replace("|", "\\|").replace("\n", "<br>")
            cells.append(value)
        rows.append(cells)
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [(r + [""] * width)[:width] for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "| " + " | ".join(["---"] * width) + " |"]
    for r in rows[1:]:
        lines.append("| " + " | ".join(r) + " |")
    return "\n".join(lines)


def _extract_images(doc, mode: str, assets_dir: Path | None) -> dict[str, str]:
    mapping: dict[str, str] = {}
    if mode == "none":
        return mapping
    try:
        rels = list(doc.part.rels.items())
    except Exception:
        try:
            rels = [(rid, rel) for rid, rel in doc.part.rels.items()]
        except Exception:
            rels = []
    index = 0
    for rid, rel in rels:
        if "image" not in rel.reltype:
            continue
        try:
            part = rel.target_part
            blob = part.blob
            ext = Path(str(part.partname)).suffix.lower() or ".png"
            if ext not in {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tiff", ".webp", ".emf", ".wmf"}:
                ext = ".png"
            index += 1
            safe_ext = ".jpg" if ext == ".jpeg" else ext
            name = f"image{index}{safe_ext}"
            if mode == "files" and assets_dir is not None:
                assets_dir.mkdir(parents=True, exist_ok=True)
                (assets_dir / name).write_bytes(blob)
                mapping[rid] = f"assets/{name}"
            elif mode == "flat" and assets_dir is not None:
                assets_dir.mkdir(parents=True, exist_ok=True)
                (assets_dir / name).write_bytes(blob)
                mapping[rid] = name
            else:
                mime = mimetypes.guess_type("x" + safe_ext)[0] or "image/png"
                b64 = base64.b64encode(blob).decode("ascii")
                mapping[rid] = f"data:{mime};base64,{b64}"
        except Exception:
            continue
    return mapping


def docx_to_markdown(doc, images_mode: str = "none", assets_dir: Path | None = None,
                     include_tables: bool = True) -> tuple[str, bool]:
    image_map = _extract_images(doc, images_mode, assets_dir)
    blocks: list[tuple[str, str]] = []
    counters: dict[int, int] = {}
    for block in _iter_blocks(doc):
        if isinstance(block, Table):
            counters.clear()
            md = _table_md(block) if include_tables else ""
            if md:
                blocks.append(("table", md))
            continue
        paragraph = block
        text = _para_md(paragraph, image_map, images_mode)
        level = _heading_level(paragraph)
        if level:
            counters.clear()
            if text:
                blocks.append(("heading", "#" * min(6, level) + " " + text))
            continue
        kind, indent = _list_kind(paragraph)
        if kind and text:
            if kind == "ol":
                counters[indent] = counters.get(indent, 0) + 1
                marker = f"{counters[indent]}."
            else:
                marker = "-"
            prefix = "  " * indent
            blocks.append(("list", f"{prefix}{marker} {text}"))
            continue
        counters.clear()
        if text:
            blocks.append(("para", text))
    merged: list[tuple[str, str]] = []
    for kind, text in blocks:
        if kind == "list" and merged and merged[-1][0] == "list":
            merged[-1] = ("list", merged[-1][1] + "\n" + text)
        else:
            merged.append((kind, text))
    md = "\n\n".join(text for _, text in merged).strip() + "\n"
    return md, bool(image_map)


def docx_to_md(job, inputs: list[Path], opts: dict):
    src = inputs[0]
    job.report(10, "读取 Word 文档")
    doc = Document(str(src))
    include_images = bool(opts.get("images", False))
    include_tables = bool(opts.get("tables", True))
    if include_images:
        staging = job.dir / "md-package"
        staging.mkdir(parents=True, exist_ok=True)
        md, has_images = docx_to_markdown(doc, "files", staging / "assets", include_tables)
        md_path = staging / f"{src.stem}.md"
        md_path.write_text(md, encoding="utf-8")
        if has_images:
            import zipfile

            dest = unique_path(job.outdir, f"{src.stem}-markdown.zip")
            with zipfile.ZipFile(dest, "w", zipfile.ZIP_DEFLATED) as zf:
                zf.write(md_path, md_path.name)
                assets = staging / "assets"
                if assets.exists():
                    for f in assets.iterdir():
                        if f.is_file():
                            zf.write(f, f"assets/{f.name}")
            job.report(90, "打包 Markdown 与图片")
            job.add_output(dest)
            return [dest]
        dest = unique_path(job.outdir, f"{src.stem}.md")
        dest.write_text(md, encoding="utf-8")
    else:
        md, _ = docx_to_markdown(doc, "none", include_tables=include_tables)
        dest = unique_path(job.outdir, f"{src.stem}.md")
        dest.write_text(md, encoding="utf-8")
    job.report(90, "写入 Markdown")
    job.add_output(dest)
    return [dest]


def docx_to_html(job, inputs: list[Path], opts: dict):
    src = inputs[0]
    job.report(15, "读取 Word 文档")
    doc = Document(str(src))
    md, _ = docx_to_markdown(doc, "datauri")
    job.report(60, "排版网页")
    html = md_to_html(md, title=src.stem)
    dest = unique_path(job.outdir, f"{src.stem}.html")
    dest.write_text(html, encoding="utf-8")
    job.report(90, "写入 HTML")
    job.add_output(dest)
    return [dest]


def _convert_word_com(source: Path, dest: Path) -> bool:
    try:
        import pythoncom
        import win32com.client
    except Exception:
        return False

    def office_pids() -> set[int]:
        try:
            wmi = win32com.client.GetObject("winmgmts:")
            query = wmi.ExecQuery(
                "SELECT ProcessId FROM Win32_Process WHERE Name='WINWORD.EXE' OR Name='wps.exe'"
            )
            return {int(proc.ProcessId) for proc in query}
        except Exception:
            return set()

    def kill_pids(pids: set[int]) -> None:
        for pid in pids:
            try:
                subprocess.run(
                    ["taskkill", "/PID", str(pid), "/F"],
                    capture_output=True, timeout=10, creationflags=0x08000000,
                )
            except Exception:
                continue

    pythoncom.CoInitialize()
    word = None
    document = None
    timer = None
    before = office_pids()
    new_pids: set[int] = set()
    try:
        for prog_id in ("Word.Application", "KWPS.Application"):
            try:
                word = win32com.client.DispatchEx(prog_id)
                break
            except Exception:
                word = None
        if word is None:
            return False
        new_pids = office_pids() - before
        timeout = max(20, int(os.environ.get("BAIXIA_WORD_TIMEOUT", "60")))
        timer = threading.Timer(timeout, lambda: kill_pids(new_pids))
        timer.daemon = True
        timer.start()
        word.Visible = False
        try:
            word.DisplayAlerts = 0
        except Exception:
            pass
        document = word.Documents.Open(
            str(source), ReadOnly=True, AddToRecentFiles=False, ConfirmConversions=False
        )
        document.ExportAsFixedFormat(str(dest), 17)
        return dest.exists() and dest.stat().st_size > 1024
    except Exception:
        try:
            dest.unlink(missing_ok=True)
        except Exception:
            pass
        return False
    finally:
        if timer is not None:
            timer.cancel()
        try:
            if document is not None:
                document.Close(False)
        except Exception:
            pass
        try:
            if word is not None:
                word.Quit()
        except Exception:
            pass
        time.sleep(0.2)
        remaining = new_pids & office_pids()
        if remaining:
            kill_pids(remaining)
        try:
            pythoncom.CoUninitialize()
        except Exception:
            pass


def _convert_libreoffice(source: Path, dest: Path) -> bool:
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        for candidate in (
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        ):
            if Path(candidate).exists():
                exe = candidate
                break
    if not exe:
        return False
    try:
        proc = subprocess.run(
            [exe, "--headless", "--convert-to", "pdf", "--outdir", str(dest.parent), str(source)],
            capture_output=True, timeout=180, creationflags=0x08000000,
        )
        produced = dest.parent / (source.stem + ".pdf")
        if proc.returncode == 0 and produced.exists():
            if produced != dest:
                produced.replace(dest)
            return dest.exists() and dest.stat().st_size > 1024
    except Exception:
        pass
    return False


def docx_to_pdf(job, inputs: list[Path], opts: dict):
    src = inputs[0]
    dest = unique_path(job.outdir, f"{src.stem}.pdf")
    job.report(15, "尝试用 Word / WPS 导出 PDF")
    if _convert_word_com(src, dest):
        job.note = "Word / WPS 排版"
        job.report(92, "Word 排版导出完成")
        job.add_output(dest)
        return [dest]
    job.report(45, "尝试用 LibreOffice 导出 PDF")
    if _convert_libreoffice(src, dest):
        job.note = "LibreOffice 排版"
        job.report(92, "LibreOffice 排版导出完成")
        job.add_output(dest)
        return [dest]
    job.report(55, "未找到 Word / LibreOffice，使用内置排版")
    doc = Document(str(src))
    assets = job.dir / "pdf-assets"
    assets.mkdir(parents=True, exist_ok=True)
    md, _ = docx_to_markdown(doc, "flat", assets)
    job.report(72, "排版 PDF")
    body = md_to_html(md, full=False)
    page_html = f"<html><body>{body}</body></html>"
    html_to_pdf(page_html, dest, archive_dir=assets)
    job.note = "内置排版（复杂版式可能失真）"
    job.report(92, "写入 PDF")
    job.add_output(dest)
    return [dest]


# ───────────────────────── markdown → docx / pdf ─────────────────────────

INLINE_TOKEN = re.compile(
    r"(\*\*\*.+?\*\*\*|\*\*.+?\*\*|__.+?__|~~.+?~~|`[^`]+`|\*[^*\n]+?\*|_[^_\n]+?_|!?\[[^\]]*\]\([^)]*\))",
    re.S,
)


def _add_hyperlink(paragraph: Paragraph, url: str, text: str) -> None:
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    try:
        r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    except Exception:
        paragraph.add_run(text)
        return
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1A5FB4")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _add_inline(paragraph: Paragraph, text: str) -> None:
    pos = 0
    for m in INLINE_TOKEN.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        token = m.group(0)
        if token.startswith("![") or (token.startswith("[") and "](" in token):
            lm = re.match(r"!?\[([^\]]*)\]\(([^)]*)\)", token)
            if lm:
                label, url = lm.group(1), lm.group(2)
                if token.startswith("!"):
                    paragraph.add_run(f"[图片: {label or 'image'}]").italic = True
                elif url:
                    _add_hyperlink(paragraph, url, label or url)
                else:
                    paragraph.add_run(label)
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif token.startswith("***"):
            run = paragraph.add_run(token[3:-3])
            run.bold = True
            run.italic = True
        elif token.startswith("**") or token.startswith("__"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("~~"):
            run = paragraph.add_run(token[2:-2])
            run.font.strike = True
        elif token.startswith("*") or token.startswith("_"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        else:
            paragraph.add_run(token)
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _style_document(doc) -> None:
    for name in ["Normal", "Body Text", "List Bullet", "List Number", "Quote", "Title"]:
        try:
            style = doc.styles[name]
            style.font.name = "DengXian"
            if style._element.rPr is not None and style._element.rPr.rFonts is not None:
                style._element.rPr.rFonts.set(qn("w:eastAsia"), "DengXian")
            if name == "Normal":
                style.font.size = Pt(10.5)
                style.paragraph_format.space_after = Pt(6)
                style.paragraph_format.line_spacing = 1.35
        except Exception:
            continue
    for i in range(1, 5):
        try:
            style = doc.styles[f"Heading {i}"]
            style.font.name = "DengXian"
            style.font.color.rgb = RGBColor(0x14, 0x14, 0x14)
            if style._element.rPr is not None and style._element.rPr.rFonts is not None:
                style._element.rPr.rFonts.set(qn("w:eastAsia"), "DengXian")
        except Exception:
            continue


def _md_blocks(md: str):
    lines = (md or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")
    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line:
            i += 1
            continue
        if line.startswith("```") or line.startswith("~~~"):
            fence = line[:3]
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith(fence):
                buf.append(lines[i])
                i += 1
            i += 1
            yield ("code", "\n".join(buf))
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            yield ("heading", (len(m.group(1)), m.group(2)))
            i += 1
            continue
        if re.match(r"^(\*\s*){3,}$|^(-\s*){3,}$|^(_\s*){3,}$", line):
            yield ("hr", "")
            i += 1
            continue
        if line.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip()[1:].lstrip())
                i += 1
            yield ("quote", " ".join(buf))
            continue
        if "|" in line and i + 1 < len(lines) and re.search(r"\|?\s*:?-{2,}", lines[i + 1]):
            table_lines = []
            while i < len(lines) and "|" in lines[i] and lines[i].strip():
                table_lines.append(lines[i].strip())
                i += 1
            yield ("table", table_lines)
            continue
        lm = re.match(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$", raw)
        if lm:
            ordered = lm.group(2)[0].isdigit()
            items = []
            while i < len(lines):
                im = re.match(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$", lines[i])
                if not im:
                    break
                items.append(im.group(3).strip())
                i += 1
            yield ("ol" if ordered else "ul", items)
            continue
        para = [line]
        i += 1
        while i < len(lines):
            nxt = lines[i]
            ns = nxt.strip()
            if (not ns or ns.startswith(("#", "```", "~~~", ">")) or re.match(r"^(\s*)([-*+]|\d+[.)])\s+", nxt)
                    or re.match(r"^(\*\s*){3,}$|^(-\s*){3,}$", ns)
                    or ("|" in ns and i + 1 < len(lines) and re.search(r"\|?\s*:?-{2,}", lines[i + 1]))):
                break
            para.append(ns)
            i += 1
        yield ("para", " ".join(para))


def md_to_docx(job, inputs: list[Path], opts: dict):
    src = inputs[0]
    job.report(10, "解析 Markdown")
    md = src.read_text(encoding="utf-8", errors="replace")
    doc = Document()
    _style_document(doc)
    for kind, payload in _md_blocks(md):
        if kind == "heading":
            level, text = payload
            paragraph = doc.add_heading("", level=min(6, max(1, level)))
            _add_inline(paragraph, text)
        elif kind == "para":
            paragraph = doc.add_paragraph()
            _add_inline(paragraph, payload)
        elif kind in {"ul", "ol"}:
            for item in payload:
                paragraph = doc.add_paragraph(style="List Bullet" if kind == "ul" else "List Number")
                _add_inline(paragraph, item)
        elif kind == "quote":
            paragraph = doc.add_paragraph(style="Quote")
            _add_inline(paragraph, payload)
        elif kind == "table":
            rows = []
            for raw in payload:
                row = raw.strip()
                if row.startswith("|"):
                    row = row[1:]
                if row.endswith("|"):
                    row = row[:-1]
                rows.append([c.strip() for c in row.split("|")])
            if len(rows) >= 2 and re.match(r"^:?-{2,}:?$", rows[1][0].replace(" ", "")):
                rows.pop(1)
            if rows:
                width = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=width)
                table.style = "Table Grid"
                for ri, row in enumerate(rows):
                    for ci in range(width):
                        cell = table.cell(ri, ci)
                        value = row[ci] if ci < len(row) else ""
                        cell.text = ""
                        _add_inline(cell.paragraphs[0], value.replace("<br>", " "))
                        if ri == 0:
                            for run in cell.paragraphs[0].runs:
                                run.bold = True
                doc.add_paragraph()
        elif kind == "code":
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(8)
            run = paragraph.add_run(payload)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "F5F5F2")
            ppr = paragraph._p.get_or_add_pPr()
            ppr.append(shd)
        elif kind == "hr":
            doc.add_paragraph()
    dest = unique_path(job.outdir, f"{src.stem}.docx")
    doc.save(dest)
    job.report(90, "写入 Word")
    job.add_output(dest)
    return [dest]


def md_to_pdf(job, inputs: list[Path], opts: dict):
    src = inputs[0]
    job.report(15, "解析 Markdown")
    md = src.read_text(encoding="utf-8", errors="replace")
    job.report(45, "排版 PDF")
    html = md_to_html(md, title=src.stem, full=False)
    dest = unique_path(job.outdir, f"{src.stem}.pdf")
    html_to_pdf(f"<html><body>{html}</body></html>", dest)
    job.report(90, "写入 PDF")
    job.add_output(dest)
    return [dest]
