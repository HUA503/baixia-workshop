from __future__ import annotations

import re
import threading
from pathlib import Path

import pymupdf
from docx import Document
from docx.shared import Pt

from .utils import open_image, parse_page_spec, unique_path

_engine = None
_engine_lock = threading.Lock()
_run_lock = threading.Lock()


def get_engine():
    global _engine
    if _engine is None:
        with _engine_lock:
            if _engine is None:
                from rapidocr_onnxruntime import RapidOCR

                _engine = RapidOCR()
    return _engine


def _normalize_result(result) -> list[dict]:
    items: list[dict] = []
    if not result:
        return items
    if isinstance(result, dict):
        result = result.get("result") or result.get("txts") or []
    for entry in result:
        try:
            box, text, score = entry[0], entry[1], entry[2]
            if not text or not str(text).strip():
                continue
            xs = [float(p[0]) for p in box]
            ys = [float(p[1]) for p in box]
            items.append({
                "text": str(text).strip(),
                "score": float(score) if score is not None else 1.0,
                "x0": min(xs), "y0": min(ys), "x1": max(xs), "y1": max(ys),
            })
        except Exception:
            continue
    return items


def _join_fragments(fragments: list[str]) -> str:
    out = ""
    for raw in fragments:
        text = raw.strip()
        if not text:
            continue
        if not out:
            out = text
        elif re.search(r"[A-Za-z0-9]$", out) and re.match(r"^[A-Za-z0-9]", text):
            out += " " + text
        else:
            out += text
    return out


def group_lines(items: list[dict]) -> list[str]:
    if not items:
        return []
    items = sorted(items, key=lambda it: ((it["y0"] + it["y1"]) / 2, it["x0"]))
    lines: list[dict] = []
    for item in items:
        cy = (item["y0"] + item["y1"]) / 2
        height = max(1.0, item["y1"] - item["y0"])
        placed = False
        for line in lines:
            same_row = abs(line["cy"] - cy) <= max(height, line["height"]) * 0.55
            gap = max(item["x0"] - line["x1"], line["x0"] - item["x1"])
            near = gap <= max(height, line["height"]) * 2.2
            if same_row and near:
                line["items"].append(item)
                line["cy"] = (line["cy"] * line["count"] + cy) / (line["count"] + 1)
                line["count"] += 1
                line["height"] = max(line["height"], height)
                line["x0"] = min(line["x0"], item["x0"])
                line["x1"] = max(line["x1"], item["x1"])
                placed = True
                break
        if not placed:
            lines.append({
                "cy": cy, "height": height, "count": 1, "items": [item],
                "x0": item["x0"], "x1": item["x1"],
            })
    lines.sort(key=lambda line: line["cy"])
    result = []
    for line in lines:
        frags = [it["text"] for it in sorted(line["items"], key=lambda it: it["x0"])]
        text = _join_fragments(frags)
        if text:
            result.append(text)
    return result


def _recognize(engine, arr) -> list[dict]:
    with _run_lock:
        result = engine(arr)[0]
    return _normalize_result(result)


def _setup_docx():
    doc = Document()
    try:
        normal = doc.styles["Normal"]
        normal.font.name = "DengXian"
        normal.font.size = Pt(10.5)
        from docx.oxml.ns import qn

        if normal._element.rPr is not None and normal._element.rPr.rFonts is not None:
            normal._element.rPr.rFonts.set(qn("w:eastAsia"), "DengXian")
    except Exception:
        pass
    return doc


def ocr_pdf(job, inputs: list[Path], opts: dict):
    import numpy as np

    src = inputs[0]
    fmt = opts.get("format", "txt")
    dpi = max(120, min(300, int(float(opts.get("dpi", 200)))))
    doc = pymupdf.open(str(src))
    if doc.needs_pass:
        doc.close()
        raise RuntimeError("这个 PDF 有打开密码，请先解密")
    total = doc.page_count
    groups = parse_page_spec(opts.get("pages", ""), total)
    pages = [p for g in groups for p in g]
    job.report(4, "加载 OCR 引擎")
    engine = get_engine()
    texts: list[tuple[int, list[str]]] = []
    out_pdf = pymupdf.open() if fmt == "searchable" else None
    try:
        for idx, pno in enumerate(pages, 1):
            job.report(8 + int((idx - 1) / len(pages) * 82), f"识别第 {idx}/{len(pages)} 页")
            page = doc[pno]
            pix = page.get_pixmap(dpi=dpi, alpha=False)
            arr = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
            if pix.n == 3:
                arr = arr[:, :, ::-1]
            elif pix.n == 4:
                arr = arr[:, :, [2, 1, 0]]
            items = _recognize(engine, arr)
            lines = group_lines(items)
            texts.append((pno + 1, lines))
            if out_pdf is not None:
                data = pix.tobytes("jpeg", jpg_quality=86)
                new_page = out_pdf.new_page(width=page.rect.width, height=page.rect.height)
                new_page.insert_image(new_page.rect, stream=data)
                sx = page.rect.width / pix.width
                sy = page.rect.height / pix.height
                for it in items:
                    font_size = max(3.0, (it["y1"] - it["y0"]) * sy * 0.78)
                    try:
                        new_page.insert_text(
                            (it["x0"] * sx, it["y1"] * sy - 0.5),
                            it["text"],
                            fontsize=font_size,
                            fontname="china-s",
                            render_mode=3,
                            color=(0, 0, 0),
                            overlay=True,
                        )
                    except Exception:
                        continue
    finally:
        doc.close()
    if fmt == "searchable":
        dest = unique_path(job.outdir, f"{src.stem}-ocr.pdf")
        job.report(92, "写入可搜索 PDF")
        out_pdf.save(str(dest), garbage=3, deflate=True)
        out_pdf.close()
        job.add_output(dest)
        return [dest]
    joined = []
    for pno, lines in texts:
        if total > 1 or len(pages) > 1:
            joined.append(f"===== 第 {pno} 页 =====")
        joined.append("\n".join(lines) if lines else "（本页未识别到文字）")
        joined.append("")
    body = "\n".join(joined).strip() + "\n"
    if fmt == "docx":
        docx_doc = _setup_docx()
        for pno, lines in texts:
            if len(pages) > 1:
                p = docx_doc.add_paragraph()
                run = p.add_run(f"第 {pno} 页")
                run.bold = True
            paragraph = docx_doc.add_paragraph()
            if lines:
                for li, line in enumerate(lines):
                    if li:
                        paragraph.add_run().add_break()
                    paragraph.add_run(line)
            else:
                paragraph.add_run("（本页未识别到文字）").italic = True
        dest = unique_path(job.outdir, f"{src.stem}-ocr.docx")
        docx_doc.save(dest)
    else:
        dest = unique_path(job.outdir, f"{src.stem}-ocr.txt")
        dest.write_text(body, encoding="utf-8")
    job.report(95, "完成")
    job.add_output(dest)
    return [dest]


def image_ocr(job, inputs: list[Path], opts: dict):
    import numpy as np

    fmt = opts.get("format", "txt")
    merge = bool(opts.get("merge", True))
    job.report(4, "加载 OCR 引擎")
    engine = get_engine()
    results: list[tuple[str, list[str]]] = []
    for idx, src in enumerate(inputs, 1):
        job.report(8 + int((idx - 1) / len(inputs) * 82), f"识别 {src.name}")
        with open_image(src) as img:
            img.load()
            rgb = img.convert("RGB")
            arr = np.array(rgb)[:, :, ::-1].copy()
            rgb.close()
        results.append((src.name, group_lines(_recognize(engine, arr))))
    outputs = []
    if merge:
        if fmt == "docx":
            doc = _setup_docx()
            for name, lines in results:
                p = doc.add_paragraph()
                p.add_run(name).bold = True
                paragraph = doc.add_paragraph()
                if lines:
                    for li, line in enumerate(lines):
                        if li:
                            paragraph.add_run().add_break()
                        paragraph.add_run(line)
                else:
                    paragraph.add_run("（未识别到文字）").italic = True
            dest = unique_path(job.outdir, f"{inputs[0].stem}-ocr.docx")
            doc.save(dest)
        else:
            parts = []
            for name, lines in results:
                parts.append(f"===== {name} =====")
                parts.append("\n".join(lines) if lines else "（未识别到文字）")
                parts.append("")
            dest = unique_path(job.outdir, f"{inputs[0].stem}-ocr.txt")
            dest.write_text("\n".join(parts).strip() + "\n", encoding="utf-8")
        outputs.append(dest)
        job.add_output(dest)
    else:
        for name, lines in results:
            stem = Path(name).stem
            if fmt == "docx":
                doc = _setup_docx()
                paragraph = doc.add_paragraph()
                for li, line in enumerate(lines):
                    if li:
                        paragraph.add_run().add_break()
                    paragraph.add_run(line)
                dest = unique_path(job.outdir, f"{stem}-ocr.docx")
                doc.save(dest)
            else:
                dest = unique_path(job.outdir, f"{stem}-ocr.txt")
                dest.write_text(("\n".join(lines) if lines else "（未识别到文字）") + "\n", encoding="utf-8")
            outputs.append(dest)
            job.add_output(dest)
    job.report(95, "完成")
    return outputs
