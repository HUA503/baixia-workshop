from __future__ import annotations

import html as html_lib
import re
from pathlib import Path

import pymupdf
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from .utils import open_image, parse_page_spec, unique_path

PDF_ERROR = "无法打开这个 PDF，文件可能已损坏或加密。若是加密文件，请先用「PDF 加解密」移除密码。"


def _open(path: Path):
    try:
        doc = pymupdf.open(str(path))
    except Exception as exc:
        raise RuntimeError(PDF_ERROR) from exc
    if doc.needs_pass:
        doc.close()
        raise RuntimeError("这个 PDF 有打开密码，请先用「PDF 加解密」解密后再处理")
    return doc


def pdf_split(job, inputs: list[Path], opts: dict):
    src = inputs[0]
    doc = _open(src)
    total = doc.page_count
    if bool(opts.get("each", False)):
        groups = [[i] for i in range(total)]
    else:
        groups = parse_page_spec(opts.get("pages", ""), total)
    outputs = []
    for i, group in enumerate(groups, 1):
        job.report(5 + int(i / len(groups) * 85), f"导出第 {i}/{len(groups)} 份")
        new = pymupdf.open()
        for p in group:
            new.insert_pdf(doc, from_page=p, to_page=p)
        if len(group) == 1:
            label = str(group[0] + 1)
        else:
            label = f"{group[0] + 1}-{group[-1] + 1}"
        dest = unique_path(job.outdir, f"{src.stem}-p{label}.pdf")
        new.save(str(dest), garbage=3, deflate=True)
        new.close()
        outputs.append(dest)
        job.add_output(dest)
    doc.close()
    return outputs


def pdf_merge(job, inputs: list[Path], opts: dict):
    merged = pymupdf.open()
    opened = []
    try:
        for i, src in enumerate(inputs, 1):
            job.report(5 + int(i / len(inputs) * 85), f"合并 {src.name}")
            doc = _open(src)
            opened.append(doc)
            merged.insert_pdf(doc)
        dest = unique_path(job.outdir, "merged.pdf")
        try:
            merged.subset_fonts()
        except Exception:
            pass
        merged.save(str(dest), garbage=3, deflate=True)
    finally:
        for d in opened:
            d.close()
        merged.close()
    job.report(95, "写入 PDF")
    job.add_output(dest)
    return [dest]


def pdf_lock(job, inputs: list[Path], opts: dict):
    src = inputs[0]
    mode = opts.get("mode", "encrypt")
    password = opts.get("password") or ""
    if not str(password).strip():
        raise ValueError("请填写密码")
    job.report(20, "读取 PDF")
    if mode == "decrypt":
        doc = pymupdf.open(str(src))
        if not doc.needs_pass:
            doc.close()
            raise RuntimeError("这个 PDF 没有打开密码，不需要解密")
        if doc.authenticate(password) <= 0:
            doc.close()
            raise RuntimeError("密码不正确")
        dest = unique_path(job.outdir, f"{src.stem}-decrypted.pdf")
        job.report(60, "移除密码")
        doc.save(str(dest), garbage=3, deflate=True)
        doc.close()
    else:
        doc = _open(src)
        dest = unique_path(job.outdir, f"{src.stem}-locked.pdf")
        job.report(60, "写入加密 PDF")
        doc.save(
            str(dest),
            encryption=pymupdf.PDF_ENCRYPT_AES_256,
            owner_pw=password + "-owner",
            user_pw=password,
            garbage=3,
            deflate=True,
        )
        doc.close()
    job.report(95, "完成")
    job.add_output(dest)
    return [dest]


def _optimize_pdf_images(doc, max_dim: int, quality: int, job, progress_start: int = 25, progress_end: int = 65) -> int:
    from io import BytesIO

    from PIL import Image

    replaced = set()
    candidates = 0
    for page in doc:
        for image in page.get_images(full=True):
            if image[0] in replaced:
                continue
            xref, smask, width, height = image[0], image[1], image[2], image[3]
            if smask:
                continue
            if not width or not height or max(width, height) < 400:
                continue
            candidates += 1
    done = 0
    for page in doc:
        for image in page.get_images(full=True):
            xref, smask, width, height = image[0], image[1], image[2], image[3]
            if xref in replaced or smask:
                continue
            try:
                info = doc.extract_image(xref)
            except Exception:
                continue
            blob = info.get("image") or b""
            ext = (info.get("ext") or "").lower()
            if len(blob) < 80_000 and max(width, height) <= max_dim:
                continue
            if ext in {"jpx", "jb2"}:
                continue
            try:
                img = Image.open(BytesIO(blob))
                img.load()
                if img.mode in {"RGBA", "LA"} or (img.mode == "P" and "transparency" in img.info):
                    continue
                if img.mode not in {"RGB", "L"}:
                    img = img.convert("RGB")
                if max(img.size) > max_dim:
                    ratio = max_dim / max(img.size)
                    img = img.resize((max(1, round(img.width * ratio)), max(1, round(img.height * ratio))), Image.Resampling.LANCZOS)
                buf = BytesIO()
                img.save(buf, "JPEG", quality=quality, optimize=True, progressive=True)
                data = buf.getvalue()
                img.close()
                if len(data) >= len(blob):
                    continue
                page.replace_image(xref, stream=data)
                replaced.add(xref)
                done += 1
                job.report(
                    progress_start + int(min(1.0, done / max(1, candidates)) * (progress_end - progress_start)),
                    f"优化内嵌图片 {done}/{candidates}",
                )
            except Exception:
                continue
    return len(replaced)


def pdf_compress(job, inputs: list[Path], opts: dict):
    src = inputs[0]
    level = opts.get("level", "light")
    job.report(10, "分析 PDF")
    doc = _open(src)
    dest = unique_path(job.outdir, f"{src.stem}-compressed.pdf")
    try:
        optimized = 0
        if level == "raster":
            job.note = "整页图像化（文字不可搜索）"
            dpi, quality = 150, 70
            out = pymupdf.open()
            total = doc.page_count
            for i, page in enumerate(doc, 1):
                job.report(10 + int(i / total * 70), f"重绘第 {i}/{total} 页")
                pix = page.get_pixmap(dpi=dpi, alpha=False)
                data = pix.tobytes("jpeg", jpg_quality=quality)
                new_page = out.new_page(width=page.rect.width, height=page.rect.height)
                new_page.insert_image(new_page.rect, stream=data)
            out.save(str(dest), garbage=4, deflate=True)
            out.close()
        else:
            if level in {"strong", "screen"}:
                max_dim, quality = 1400, 66
            elif level in {"balanced", "ebook", "print"}:
                max_dim, quality = 2200, 82
            else:
                max_dim, quality = 0, 0
            if max_dim:
                optimized = _optimize_pdf_images(doc, max_dim, quality, job)
                job.note = f"优化 {optimized} 张内嵌图片 · 文字保留"
                job.report(68, f"已优化 {optimized} 张内嵌图片 · 精简字体")
            else:
                job.note = "无损清理 · 文字保留"
                job.report(50, "精简内嵌字体")
            try:
                doc.subset_fonts()
            except Exception:
                pass
            job.report(72, "清理冗余数据")
            doc.save(str(dest), garbage=4, deflate=True, clean=True)
    finally:
        doc.close()
    original = src.stat().st_size
    if dest.exists() and dest.stat().st_size >= original:
        dest.unlink()
        import shutil

        shutil.copy2(src, dest)
        job.message = "压缩后体积反而变大，已保留原文件"
        job.note = "已保留原文件（压缩无收益）"
    job.report(95, "完成")
    job.add_output(dest)
    return [dest]


def pdf_to_images(job, inputs: list[Path], opts: dict):
    src = inputs[0]
    fmt = "jpg" if opts.get("format") == "jpg" else "png"
    dpi = max(72, min(400, int(float(opts.get("dpi", 150)))))
    doc = _open(src)
    groups = parse_page_spec(opts.get("pages", ""), doc.page_count)
    pages = [p for g in groups for p in g]
    outputs = []
    for i, pno in enumerate(pages, 1):
        job.report(5 + int(i / len(pages) * 85), f"渲染第 {i}/{len(pages)} 页")
        page = doc[pno]
        pix = page.get_pixmap(dpi=dpi, alpha=False)
        dest = unique_path(job.outdir, f"{src.stem}-p{pno + 1:03d}.{fmt}")
        if fmt == "jpg":
            pix.save(str(dest), jpg_quality=90)
        else:
            pix.save(str(dest))
        outputs.append(dest)
        job.add_output(dest)
    doc.close()
    return outputs


def images_to_pdf(job, inputs: list[Path], opts: dict):
    page_size = opts.get("page_size", "a4")
    margin = max(0, min(40, float(opts.get("margin", 12) or 0))) * 72 / 25.4
    out = pymupdf.open()
    temp_dir = job.dir / "tmp-img"
    temp_dir.mkdir(parents=True, exist_ok=True)
    for i, src in enumerate(inputs, 1):
        job.report(5 + int(i / len(inputs) * 80), f"排版 {i}/{len(inputs)}")
        with open_image(src) as img:
            img.load()
            rgba = img.convert("RGBA")
            norm = temp_dir / f"{i:04d}.png"
            rgba.save(norm, "PNG")
            iw, ih = rgba.size
        if page_size == "auto":
            width_pt = iw * 72 / 96
            height_pt = ih * 72 / 96
            page = out.new_page(width=width_pt, height=height_pt)
            page.insert_image(page.rect, filename=str(norm))
        else:
            page = out.new_page(width=595.28, height=841.89)
            box = pymupdf.Rect(margin, margin, 595.28 - margin, 841.89 - margin)
            ratio = min(box.width / iw, box.height / ih)
            dw, dh = iw * ratio, ih * ratio
            target = pymupdf.Rect(0, 0, dw, dh)
            target.x0 = box.x0 + (box.width - dw) / 2
            target.y0 = box.y0 + (box.height - dh) / 2
            target.x1 = target.x0 + dw
            target.y1 = target.y0 + dh
            page.insert_image(target, filename=str(norm))
    dest = unique_path(job.outdir, f"{inputs[0].stem}-images.pdf" if len(inputs) == 1 else "images.pdf")
    job.report(90, "写入 PDF")
    out.save(str(dest), garbage=3, deflate=True)
    out.close()
    job.add_output(dest)
    return [dest]


# ───────────────────────── PDF text extraction ─────────────────────────

def _weighted_median_size(doc) -> float:
    weights: dict[float, int] = {}
    sample = range(min(doc.page_count, 30))
    for idx in sample:
        try:
            data = doc[idx].get_text("dict")
        except Exception:
            continue
        for block in data.get("blocks", []):
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    text = (span.get("text") or "").strip()
                    if not text:
                        continue
                    size = round(float(span.get("size") or 10), 1)
                    weights[size] = weights.get(size, 0) + len(text)
    if not weights:
        return 10.0
    ordered = sorted(weights.items())
    total = sum(weights.values())
    acc = 0
    for size, count in ordered:
        acc += count
        if acc >= total / 2:
            return size
    return ordered[-1][0]


def _smart_join(parts: list[str]) -> str:
    out = ""
    for raw in parts:
        text = raw.strip()
        if not text:
            continue
        if not out:
            out = text
        elif out.endswith("-") and text[:1].islower():
            out = out[:-1] + text
        elif re.search(r"[A-Za-z0-9,.;:!?)]$", out) and re.match(r"^[A-Za-z0-9(]", text):
            out += " " + text
        else:
            out += text
    return out


BULLET_RE = re.compile(r"^[•·▪◦●○◆◇‣]\s*")
ORDERED_RE = re.compile(r"^(\d{1,3})[.)、]\s+")


def _table_rows(table) -> list[list[str]]:
    rows = []
    for row in table.extract():
        cleaned = []
        for cell in row:
            cleaned.append(re.sub(r"\s*\n\s*", " ", (cell or "").strip()))
        rows.append(cleaned)
    return rows


def _md_table_from_rows(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [(r + [""] * width)[:width] for r in rows]
    escaped = [[c.replace("|", "\\|") for c in r] for r in rows]
    lines = ["| " + " | ".join(escaped[0]) + " |", "| " + " | ".join(["---"] * width) + " |"]
    for r in escaped[1:]:
        lines.append("| " + " | ".join(r) + " |")
    return "\n".join(lines)


def _docx_setup(doc) -> None:
    try:
        normal = doc.styles["Normal"]
        normal.font.name = "DengXian"
        normal.font.size = Pt(10.5)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.35
        if normal._element.rPr is not None and normal._element.rPr.rFonts is not None:
            normal._element.rPr.rFonts.set(qn("w:eastAsia"), "DengXian")
        for i in range(1, 5):
            style = doc.styles[f"Heading {i}"]
            style.font.name = "DengXian"
            style.font.color.rgb = RGBColor(0x14, 0x14, 0x14)
            if style._element.rPr is not None and style._element.rPr.rFonts is not None:
                style._element.rPr.rFonts.set(qn("w:eastAsia"), "DengXian")
    except Exception:
        pass


def _page_items(page, median: float, include_tables: bool):
    table_rects = []
    tables = []
    if include_tables:
        try:
            finder = page.find_tables()
            for tab in finder.tables:
                rows = _table_rows(tab)
                if rows and any(any(c for c in row) for row in rows):
                    tables.append((tab.bbox, rows))
                    table_rects.append(pymupdf.Rect(tab.bbox))
        except Exception:
            pass
    data = page.get_text("dict")
    items = []
    for bbox, rows in tables:
        items.append((bbox[1], "table", rows))
    for block in data.get("blocks", []):
        if block.get("type") != 0:
            continue
        bbox = block.get("bbox") or (0, 0, 0, 0)
        center = pymupdf.Point((bbox[0] + bbox[2]) / 2, (bbox[1] + bbox[3]) / 2)
        if any(r.contains(center) for r in table_rects):
            continue
        lines = []
        max_size = 0.0
        for line in block.get("lines", []):
            text = _smart_join([span.get("text", "") for span in line.get("spans", [])])
            if text.strip():
                lines.append(text.strip())
                for span in line.get("spans", []):
                    max_size = max(max_size, float(span.get("size") or 0))
        text = _smart_join(lines)
        if text:
            items.append((bbox[1], "text", (text, max_size)))
    items.sort(key=lambda x: x[0])
    return items


def _heading_level(size: float, median: float) -> int:
    if median <= 0 or size <= 0:
        return 0
    ratio = size / median
    if ratio >= 1.8:
        return 1
    if ratio >= 1.35:
        return 2
    if ratio >= 1.12:
        return 3
    return 0


def pdf_extract(job, inputs: list[Path], opts: dict):
    src = inputs[0]
    fmt = opts.get("format", "docx")
    include_tables = bool(opts.get("tables", True))
    markers = bool(opts.get("page_markers", False))
    job.report(8, "分析文档结构")
    doc = _open(src)
    median = _weighted_median_size(doc)
    docx_doc = None
    if fmt == "docx":
        docx_doc = Document()
        _docx_setup(docx_doc)
    md_parts: list[str] = []
    html_parts: list[str] = []
    txt_parts: list[str] = []
    any_text = False
    for pno, page in enumerate(doc, 1):
        job.report(10 + int(pno / doc.page_count * 80), f"提取第 {pno}/{doc.page_count} 页")
        if fmt == "txt":
            text = page.get_text("text")
            if text.strip():
                any_text = True
            if markers:
                txt_parts.append(f"===== 第 {pno} 页 =====\n")
            txt_parts.append(text.rstrip() + "\n")
            continue
        if fmt == "html":
            frag = page.get_text("html")
            if re.sub(r"<[^>]+>", "", frag).strip():
                any_text = True
            html_parts.append(f'<section class="pdf-page" id="page-{pno}">{frag}</section>')
            continue
        items = _page_items(page, median, include_tables)
        if fmt == "md" and markers:
            md_parts.append(f"---\n\n**第 {pno} 页**\n")
        if fmt == "docx" and markers:
            docx_doc.add_paragraph(f"— 第 {pno} 页 —").runs[0].italic = True
        for _, kind, payload in items:
            if kind == "table":
                any_text = True
                rows = payload
                if fmt == "docx":
                    width = max(len(r) for r in rows)
                    table = docx_doc.add_table(rows=len(rows), cols=width)
                    table.style = "Table Grid"
                    for ri, row in enumerate(rows):
                        for ci in range(width):
                            cell = table.cell(ri, ci)
                            cell.text = ""
                            run = cell.paragraphs[0].add_run(row[ci] if ci < len(row) else "")
                            if ri == 0:
                                run.bold = True
                    docx_doc.add_paragraph()
                else:
                    md_parts.append(_md_table_from_rows(rows) + "\n")
                continue
            text, size = payload
            any_text = True
            level = _heading_level(size, median)
            bullet = BULLET_RE.match(text)
            ordered = ORDERED_RE.match(text)
            if fmt == "md":
                if level and len(text) <= 90:
                    md_parts.append("#" * level + " " + text + "\n")
                elif bullet:
                    md_parts.append("- " + BULLET_RE.sub("", text) + "\n")
                elif ordered:
                    md_parts.append(ORDERED_RE.sub(lambda m: f"{m.group(1)}. ", text) + "\n")
                else:
                    md_parts.append(text + "\n")
            else:
                if level and len(text) <= 90:
                    p = docx_doc.add_heading("", level=min(4, level))
                    p.add_run(text)
                elif bullet:
                    p = docx_doc.add_paragraph(style="List Bullet")
                    p.add_run(BULLET_RE.sub("", text))
                elif ordered:
                    p = docx_doc.add_paragraph(style="List Number")
                    p.add_run(ORDERED_RE.sub("", text))
                else:
                    p = docx_doc.add_paragraph()
                    p.add_run(text)
    doc.close()
    if not any_text:
        note = "未检测到文字层。这个 PDF 可能是扫描件，请改用「PDF OCR 识别」。"
        if fmt == "docx":
            docx_doc.add_paragraph(note).runs[0].italic = True
        elif fmt == "md":
            md_parts.append(f"> {note}\n")
        elif fmt == "html":
            html_parts.append(f"<p>{html_lib.escape(note)}</p>")
        else:
            txt_parts.append(note + "\n")
    if fmt == "docx":
        dest = unique_path(job.outdir, f"{src.stem}.docx")
        docx_doc.save(dest)
    elif fmt == "md":
        dest = unique_path(job.outdir, f"{src.stem}.md")
        dest.write_text("\n".join(md_parts), encoding="utf-8")
    elif fmt == "html":
        body = "\n".join(html_parts)
        css = """
        body{margin:0;background:#f5f5f2;color:#141414;font-family:-apple-system,'Segoe UI','Microsoft YaHei',sans-serif;}
        .wrap{max-width:900px;margin:0 auto;padding:40px 20px;}
        .pdf-page{background:#fff;border:1px solid #e6e6e0;padding:32px;margin:0 0 24px;overflow:auto;}
        p{margin:0 0 6px;line-height:1.6;}
        """
        dest = unique_path(job.outdir, f"{src.stem}.html")
        dest.write_text(
            "<!DOCTYPE html><html lang=\"zh-CN\"><head><meta charset=\"utf-8\">"
            f"<title>{html_lib.escape(src.stem)}</title><style>{css}</style></head>"
            f"<body><div class=\"wrap\">{body}</div></body></html>",
            encoding="utf-8",
        )
    else:
        dest = unique_path(job.outdir, f"{src.stem}.txt")
        dest.write_text("\n".join(txt_parts), encoding="utf-8")
    job.report(95, "完成")
    job.add_output(dest)
    return [dest]


def pdf_pages(job, inputs: list[Path], opts: dict):
    import json

    src = inputs[0]
    doc = _open(src)
    total = doc.page_count
    raw = opts.get("page_ops") or []
    if isinstance(raw, str):
        try:
            raw = json.loads(raw)
        except ValueError as exc:
            doc.close()
            raise ValueError("页面顺序数据无效") from exc
    if isinstance(raw, dict):
        raw = raw.get("pages", [])
    if not isinstance(raw, list) or not raw:
        doc.close()
        raise ValueError("请至少保留一页")
    operations = []
    for item in raw:
        try:
            source_index = int(item.get("src"))
            rotation = int(item.get("rotate", 0)) % 360
        except (TypeError, ValueError, AttributeError) as exc:
            doc.close()
            raise ValueError("页面操作数据无效") from exc
        if source_index < 0 or source_index >= total:
            doc.close()
            raise ValueError("页面序号超出范围")
        operations.append((source_index, rotation))
    if len(operations) > total * 4:
        doc.close()
        raise ValueError("页面数量过多")
    job.report(18, "重排 PDF 页面")
    new = pymupdf.open()
    try:
        for source_index, rotation in operations:
            new.insert_pdf(doc, from_page=source_index, to_page=source_index)
            page = new[-1]
            if rotation:
                page.set_rotation((page.rotation + rotation) % 360)
        dest = unique_path(job.outdir, f"{src.stem}-arranged.pdf")
        job.report(78, "写入 PDF")
        try:
            new.subset_fonts()
        except Exception:
            pass
        new.save(str(dest), garbage=3, deflate=True)
    finally:
        new.close()
        doc.close()
    job.note = f"{len(operations)} 页"
    job.report(96, "完成")
    job.add_output(dest)
    return [dest]


def _watermark_tile(text: str, font, color):
    from PIL import Image, ImageDraw

    bbox = font.getbbox(text)
    pad = max(10, int(getattr(font, "size", 24) * 0.35))
    width = max(1, bbox[2] - bbox[0] + pad * 2)
    height = max(1, bbox[3] - bbox[1] + pad * 2)
    tile = Image.new("RGBA", (width, height), (0, 0, 0, 0))
    draw = ImageDraw.Draw(tile)
    draw.text((pad - bbox[0], pad - bbox[1]), text, font=font, fill=color)
    return tile


def _watermark_overlay(page, text: str, font_size: int, angle: int, color, opacity: float, tile: bool):
    from io import BytesIO

    from PIL import Image, ImageFont

    rect = page.rect
    scale = min(2.0, 3000 / max(1.0, rect.width, rect.height))
    width = max(1, int(rect.width * scale))
    height = max(1, int(rect.height * scale))
    canvas = Image.new("RGBA", (width, height), (0, 0, 0, 0))
    size = max(8, int(font_size * scale))
    font = None
    for candidate in (
        r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\msyhbd.ttc",
        r"C:\Windows\Fonts\simhei.ttf",
        r"C:\Windows\Fonts\simsun.ttc",
    ):
        try:
            font = ImageFont.truetype(candidate, size)
            break
        except Exception:
            continue
    if font is None:
        font = ImageFont.load_default()
        text = text.encode("ascii", "replace").decode("ascii") or "WATERMARK"
    rgba = (*color, max(8, min(255, int(255 * opacity))))
    stamp = _watermark_tile(text, font, rgba)
    if angle:
        stamp = stamp.rotate(angle, expand=True, resample=Image.Resampling.BICUBIC)
    if tile:
        step_x = max(30, int(stamp.width * 0.90))
        step_y = max(30, int(stamp.height * 0.95))
        row = 0
        for y in range(-stamp.height, height + stamp.height, step_y):
            offset = (row % 2) * (step_x // 2)
            for x in range(-stamp.width, width + stamp.width, step_x):
                canvas.alpha_composite(stamp, (x + offset, y))
            row += 1
    else:
        canvas.alpha_composite(stamp, ((width - stamp.width) // 2, (height - stamp.height) // 2))
    buf = BytesIO()
    canvas.save(buf, "PNG", optimize=True)
    return buf.getvalue()


def pdf_watermark(job, inputs: list[Path], opts: dict):
    src = inputs[0]
    text = " ".join(str(opts.get("text", "") or "").split())
    if not text:
        raise ValueError("请填写水印文字")
    if len(text) > 120:
        raise ValueError("水印文字不能超过 120 个字符")
    opacity = max(0.05, min(1.0, float(opts.get("opacity", 18)) / 100.0))
    font_size = max(8, min(160, int(float(opts.get("font_size", 36)))))
    angle = max(-90, min(90, int(float(opts.get("angle", 30)))))
    tile = bool(opts.get("tile", True))
    palette = {
        "gray": (110, 110, 110),
        "red": (190, 42, 42),
        "blue": (30, 86, 170),
        "black": (25, 25, 25),
        "green": (35, 125, 80),
    }
    color = palette.get(str(opts.get("color", "gray")), palette["gray"])
    doc = _open(src)
    total = doc.page_count
    groups = parse_page_spec(opts.get("pages", ""), total)
    selected = []
    for group in groups:
        selected.extend(group)
    seen = set()
    selected = [i for i in selected if not (i in seen or seen.add(i))]
    for index, page_no in enumerate(selected, 1):
        job.report(8 + int(index / max(1, len(selected)) * 84), f"添加水印 {index}/{len(selected)}")
        page = doc[page_no]
        stream = _watermark_overlay(page, text, font_size, angle, color, opacity, tile)
        page.insert_image(page.rect, stream=stream, overlay=True)
    dest = unique_path(job.outdir, f"{src.stem}-watermark.pdf")
    job.report(94, "写入 PDF")
    try:
        doc.subset_fonts()
    except Exception:
        pass
    doc.save(str(dest), garbage=3, deflate=True)
    doc.close()
    job.note = f"{len(selected)} 页 · 文字水印"
    job.report(97, "完成")
    job.add_output(dest)
    return [dest]
